"""사용자 Excel 머지 라운드트립 테스트.

목적: 셀 색상, 수식, 병합셀이 머지 후에도 보존되는지 확인.
"""
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.models.compare_session import DiffType
from app.schemas.user_sheet import ColumnMapping
from app.services.merge_engine import apply_merge


def _make_user_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "단가표"
    ws.append(["코드", "모델명", "카테고리", "정가", "딜러가"])
    ws.append(["2767-20", "임팩트", "공구", 500000, 350000])
    ws.append(["2780-20", "그라인더", "공구", 300000, 220000])
    # 색상 + 굵게
    ws["A2"].font = Font(bold=True)
    ws["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    # 수식
    ws["F1"] = "마진"
    ws["F2"] = "=D2-E2"
    ws["F3"] = "=D3-E3"
    # 병합셀 (헤더 위에 메모)
    ws.merge_cells("H1:I1")
    ws["H1"] = "메모"
    wb.save(path)


def test_merge_preserves_styles_formulas_and_merges(tmp_path: Path):
    src = tmp_path / "user.xlsx"
    out = tmp_path / "user_merged.xlsx"
    _make_user_workbook(src)

    mapping = ColumnMapping(
        sku="A",
        model_name="B",
        category="C",
        list_price="D",
        dealer_price="E",
    )

    selected_diffs = [
        # 가격 변경: 2767-20의 정가만 변경
        {
            "diff_type": DiffType.PRICE_CHANGED.value,
            "sku": "2767-20",
            "before": {"list_price": "500000", "dealer_price": "350000"},
            "after": {"list_price": "520000", "dealer_price": "350000"},
            "source_ref": {"row": 2},
        },
        # 신규 추가
        {
            "diff_type": DiffType.ADDED.value,
            "sku": "2855-20",
            "before": None,
            "after": {
                "model_name": "신제품 드릴",
                "category": "공구",
                "list_price": "400000",
                "dealer_price": "280000",
            },
            "source_ref": None,
        },
    ]

    result = apply_merge(
        src_xlsx=src,
        out_xlsx=out,
        sheet_name="단가표",
        mapping=mapping,
        selected_diffs=selected_diffs,
        add_new_rows=True,
        mark_removed=True,
    )
    assert result.applied == 2

    wb = load_workbook(out)
    ws = wb["단가표"]

    # 가격 변경 적용
    assert ws["D2"].value == 520000.0

    # 스타일 보존(굵게 + 노란 배경)
    assert ws["A2"].font.bold is True
    assert ws["A2"].fill.fgColor.value.upper().endswith("FFFF00")

    # 수식 보존
    assert str(ws["F2"].value).startswith("=")
    assert str(ws["F3"].value).startswith("=")

    # 병합셀 보존
    assert "H1:I1" in [str(r) for r in ws.merged_cells.ranges]

    # 신규 행 삽입 (4행)
    assert ws["A4"].value == "2855-20"
    assert ws["B4"].value == "신제품 드릴"
    assert float(ws["D4"].value) == 400000.0


def test_merge_records_removed_in_separate_sheet(tmp_path: Path):
    src = tmp_path / "user2.xlsx"
    out = tmp_path / "user2_merged.xlsx"
    _make_user_workbook(src)

    mapping = ColumnMapping(sku="A", model_name="B", list_price="D", dealer_price="E")
    diffs = [
        {
            "diff_type": DiffType.REMOVED.value,
            "sku": "2780-20",
            "before": {
                "model_name": "그라인더",
                "list_price": "300000",
                "dealer_price": "220000",
                "category": "공구",
            },
            "after": None,
            "source_ref": {"row": 3},
        }
    ]
    apply_merge(
        src_xlsx=src,
        out_xlsx=out,
        sheet_name="단가표",
        mapping=mapping,
        selected_diffs=diffs,
        mark_removed=True,
    )
    wb = load_workbook(out)
    assert "_단종후보" in wb.sheetnames
    discontinued = wb["_단종후보"]
    rows = list(discontinued.iter_rows(values_only=True))
    # 헤더 + 1행
    assert len(rows) >= 2
    assert "2780-20" in [r[0] for r in rows[1:]]
    # 원본 시트의 해당 행은 그대로 유지
    main = wb["단가표"]
    assert main["A3"].value == "2780-20"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
