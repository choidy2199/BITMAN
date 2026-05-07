"""Excel 읽기 헬퍼. openpyxl로 시트/헤더/행을 추출."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.schemas.user_sheet import ColumnMapping
from app.services.normalize import (
    col_index,
    col_letter,
    infer_column_mapping,
    normalize_sku,
    normalize_text,
    parse_price,
)


def preview_workbook(path: str | Path, max_sheets: int = 5) -> dict:
    """업로드 직후 시트/헤더/행수 추정 + 매핑 자동 추론."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets_info = []
    suggested = None
    for ws in wb.worksheets[:max_sheets]:
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [(col_letter(i), str(v) if v is not None else "") for i, v in enumerate(first_row)]
        # data_rows 카운트(헤더 제외)
        data_rows = max(0, ws.max_row - 1) if ws.max_row else 0
        sheets_info.append(
            {
                "name": ws.title,
                "headers": [{"col": c, "value": v} for c, v in headers],
                "data_rows": data_rows,
            }
        )
        if suggested is None and headers:
            mapping = infer_column_mapping([v for _, v in headers])
            if "sku" in mapping:
                # mapping은 0-based index 문자열 → Excel 컬럼 문자로 변환
                converted = {k: col_letter(int(v)) for k, v in mapping.items()}
                suggested = ColumnMapping(
                    sku=converted.get("sku", "A"),
                    model_name=converted.get("model_name"),
                    list_price=converted.get("list_price"),
                    dealer_price=converted.get("dealer_price"),
                    category=converted.get("category"),
                )
    wb.close()
    return {"sheets": sheets_info, "suggested_mapping": suggested}


def read_rows(
    path: str | Path,
    sheet_name: str,
    header_row: int,
    mapping: ColumnMapping,
) -> list[dict[str, Any]]:
    """매핑에 따라 행을 표준 dict로 추출.

    반환 형식: [{"sku": ..., "model_name": ..., "list_price": Decimal,
                 "dealer_price": Decimal, "category": ..., "_row": <int>}]
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet not found: {sheet_name}")
    ws = wb[sheet_name]

    sku_idx = col_index(mapping.sku)
    model_idx = col_index(mapping.model_name) if mapping.model_name else None
    list_idx = col_index(mapping.list_price) if mapping.list_price else None
    dealer_idx = col_index(mapping.dealer_price) if mapping.dealer_price else None
    cat_idx = col_index(mapping.category) if mapping.category else None

    out: list[dict[str, Any]] = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        if not row:
            continue
        sku = normalize_sku(row[sku_idx] if sku_idx < len(row) else None)
        if not sku:
            continue
        out.append(
            {
                "sku": sku,
                "model_name": normalize_text(row[model_idx]) if model_idx is not None and model_idx < len(row) else None,
                "list_price": parse_price(row[list_idx]) if list_idx is not None and list_idx < len(row) else None,
                "dealer_price": parse_price(row[dealer_idx]) if dealer_idx is not None and dealer_idx < len(row) else None,
                "category": normalize_text(row[cat_idx]) if cat_idx is not None and cat_idx < len(row) else None,
                "_row": row_idx,
            }
        )
    wb.close()
    return out
