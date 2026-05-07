"""사용자 Excel 포맷 보존 머지.

핵심 원칙:
- 원본을 다시 만들지 말고, openpyxl로 열어 셀 단위로 값만 갱신.
- 신규 행은 마지막 데이터 행 아래에 삽입하고 윗 행 스타일을 복제.
- 단종(REMOVED)은 기본적으로 별도 시트(`_단종후보`)에 기록하고 원본 행은 보존.
- 수식/병합셀/스타일은 직접 변경하지 않음.
"""
from __future__ import annotations

import shutil
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from app.models.compare_session import DiffType
from app.schemas.user_sheet import ColumnMapping
from app.services.normalize import col_index


@dataclass
class MergeResult:
    output_path: Path
    warnings: list[str] = field(default_factory=list)
    applied: int = 0
    skipped: int = 0


def _coerce_value(raw: Any) -> Any:
    """JSON-serialized Decimal(str) 등을 Excel에 적합한 값으로 변환."""
    if raw is None:
        return None
    if isinstance(raw, str):
        # 가격일 가능성 — 숫자로 변환 시도
        try:
            return float(Decimal(raw))
        except Exception:
            return raw
    return raw


def _set_cell_value(ws, row: int, col_letter: str, value: Any) -> None:
    cell = ws[f"{col_letter}{row}"]
    cell.value = value


def _last_data_row(ws, sku_col_idx: int) -> int:
    """SKU 컬럼이 비어있지 않은 마지막 행."""
    last = 1
    for r_idx in range(ws.max_row, 0, -1):
        cell = ws.cell(row=r_idx, column=sku_col_idx + 1)
        if cell.value not in (None, ""):
            last = r_idx
            break
    return last


def _copy_row_style(ws, src_row: int, dst_row: int) -> None:
    for cell in ws[src_row]:
        new_cell = ws.cell(row=dst_row, column=cell.column)
        if cell.has_style:
            new_cell.font = copy(cell.font)
            new_cell.border = copy(cell.border)
            new_cell.fill = copy(cell.fill)
            new_cell.number_format = cell.number_format
            new_cell.alignment = copy(cell.alignment)
            new_cell.protection = copy(cell.protection)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            new_cell.value = Translator(cell.value, origin=cell.coordinate).translate_formula(
                f"{get_column_letter(cell.column)}{dst_row}"
            )


def _ensure_removed_sheet(wb):
    name = "_단종후보"
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    ws.append(["SKU", "모델명", "정가", "딜러가", "카테고리", "기록일"])
    return ws


def apply_merge(
    src_xlsx: str | Path,
    out_xlsx: str | Path,
    sheet_name: str,
    mapping: ColumnMapping,
    selected_diffs: list[dict[str, Any]],
    *,
    add_new_rows: bool = True,
    mark_removed: bool = True,
) -> MergeResult:
    """`src_xlsx`를 `out_xlsx`로 복사 후 선택된 diff를 셀 단위 적용.

    selected_diffs: 각 항목은 {"diff_type": "...", "sku": "...", "before": {...},
                              "after": {...}, "source_ref": {"row": int}}
    """
    src_path = Path(src_xlsx)
    out_path = Path(out_xlsx)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src_path, out_path)

    wb = load_workbook(out_path)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet not found in workbook: {sheet_name}")
    ws = wb[sheet_name]

    sku_idx = col_index(mapping.sku)
    model_letter = mapping.model_name
    list_letter = mapping.list_price
    dealer_letter = mapping.dealer_price
    cat_letter = mapping.category

    warnings: list[str] = []
    applied = 0
    skipped = 0
    next_new_row = _last_data_row(ws, sku_idx) + 1
    style_template_row = _last_data_row(ws, sku_idx) or 1

    removed_sheet = _ensure_removed_sheet(wb) if mark_removed else None

    for d in selected_diffs:
        dtype = d.get("diff_type")
        sku = d.get("sku")
        after = d.get("after") or {}
        before = d.get("before") or {}
        ref = d.get("source_ref") or {}

        if dtype in (
            DiffType.PRICE_CHANGED.value,
            DiffType.MODEL_NAME_CHANGED.value,
            DiffType.CATEGORY_CHANGED.value,
        ):
            row = ref.get("row")
            if not row:
                warnings.append(f"{sku}: source_ref.row 누락, 건너뜀")
                skipped += 1
                continue
            if model_letter and after.get("model_name") is not None:
                _set_cell_value(ws, row, model_letter, after.get("model_name"))
            if list_letter and after.get("list_price") is not None:
                _set_cell_value(ws, row, list_letter, _coerce_value(after.get("list_price")))
            if dealer_letter and after.get("dealer_price") is not None:
                _set_cell_value(ws, row, dealer_letter, _coerce_value(after.get("dealer_price")))
            if cat_letter and after.get("category") is not None:
                _set_cell_value(ws, row, cat_letter, after.get("category"))
            applied += 1

        elif dtype == DiffType.ADDED.value:
            if not add_new_rows:
                skipped += 1
                continue
            target_row = next_new_row
            _copy_row_style(ws, style_template_row, target_row)
            _set_cell_value(ws, target_row, mapping.sku, sku)
            if model_letter:
                _set_cell_value(ws, target_row, model_letter, after.get("model_name"))
            if list_letter and after.get("list_price") is not None:
                _set_cell_value(ws, target_row, list_letter, _coerce_value(after.get("list_price")))
            if dealer_letter and after.get("dealer_price") is not None:
                _set_cell_value(
                    ws, target_row, dealer_letter, _coerce_value(after.get("dealer_price"))
                )
            if cat_letter:
                _set_cell_value(ws, target_row, cat_letter, after.get("category"))
            next_new_row += 1
            applied += 1

        elif dtype == DiffType.REMOVED.value:
            if removed_sheet is not None:
                removed_sheet.append(
                    [
                        sku,
                        before.get("model_name"),
                        _coerce_value(before.get("list_price")),
                        _coerce_value(before.get("dealer_price")),
                        before.get("category"),
                        "",
                    ]
                )
                applied += 1
            else:
                skipped += 1

        elif dtype == DiffType.SKU_CHANGED_SUSPECTED.value:
            warnings.append(f"{sku}: SKU 변경 의심은 사용자 확정 후 적용 필요, 건너뜀")
            skipped += 1
        else:
            skipped += 1

    wb.save(out_path)
    wb.close()
    return MergeResult(output_path=out_path, warnings=warnings, applied=applied, skipped=skipped)
